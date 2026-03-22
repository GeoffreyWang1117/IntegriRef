#!/usr/bin/env python3
"""Generate IntegriRef project introduction slides (Chinese, casual tone)."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── Color palette ──
BG_DARK = RGBColor(0x1A, 0x1A, 0x2E)       # deep navy
BG_CARD = RGBColor(0x16, 0x21, 0x3E)       # card background
ACCENT  = RGBColor(0x00, 0xD2, 0xFF)       # cyan accent
ACCENT2 = RGBColor(0x7C, 0x3A, 0xED)       # purple accent
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
GRAY    = RGBColor(0xBB, 0xBB, 0xBB)
GREEN   = RGBColor(0x10, 0xB9, 0x81)
RED     = RGBColor(0xEF, 0x44, 0x44)
YELLOW  = RGBColor(0xFB, 0xBF, 0x24)
ORANGE  = RGBColor(0xF9, 0x73, 0x16)


def set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text_box(slide, left, top, width, height, text, font_size=18,
                 color=WHITE, bold=False, alignment=PP_ALIGN.LEFT,
                 font_name="Microsoft YaHei"):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                     Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox, tf


def add_para(tf, text, font_size=16, color=WHITE, bold=False, space_before=6,
             alignment=PP_ALIGN.LEFT, font_name="Microsoft YaHei"):
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    p.space_before = Pt(space_before)
    return p


def add_rounded_rect(slide, left, top, width, height, fill_color):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_card(slide, left, top, width, height, title, bullets,
             title_color=ACCENT, bullet_color=GRAY):
    add_rounded_rect(slide, left, top, width, height, BG_CARD)
    _, tf = add_text_box(slide, left + 0.15, top + 0.1, width - 0.3, 0.4,
                         title, font_size=14, color=title_color, bold=True)
    for b in bullets:
        add_para(tf, b, font_size=12, color=bullet_color, space_before=4)


# ════════════════════════════════════════
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# ════════════════════════════════════════
# SLIDE 1 — Title
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
set_slide_bg(slide, BG_DARK)

add_text_box(slide, 1.5, 1.5, 10, 1.2,
             "IntegriRef", font_size=54, color=ACCENT, bold=True)
add_text_box(slide, 1.5, 2.7, 10, 0.8,
             "一个帮你检查「参考文献到底靠不靠谱」的工具",
             font_size=24, color=WHITE)

# subtitle bullets
_, tf = add_text_box(slide, 1.5, 3.8, 10, 2.5, "", font_size=18)
items = [
    "你引用的论文真的存在吗？DOI 是不是编的？",
    "引用说「支持」，原文其实说的是「反对」？",
    "有没有引了已经被撤稿的论文？",
    "ChatGPT 帮你写的参考文献，有多少是幻觉？",
]
for item in items:
    add_para(tf, f"  {item}", font_size=18, color=GRAY, space_before=10)

add_text_box(slide, 1.5, 6.2, 10, 0.5,
             "GitHub: GeoffreyWang1117/IntegriRef",
             font_size=14, color=GRAY)

# ════════════════════════════════════════
# SLIDE 2 — Problem: Why do we need this?
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_DARK)

add_text_box(slide, 0.8, 0.4, 12, 0.8,
             "我们到底在解决什么问题？", font_size=36, color=WHITE, bold=True)

# Left column - problems
add_rounded_rect(slide, 0.8, 1.5, 5.5, 5.2, BG_CARD)
_, tf = add_text_box(slide, 1.0, 1.6, 5.1, 5.0,
                     "学术引用里的「坑」", font_size=20, color=RED, bold=True)
problems = [
    ("假引用", "LLM 生成的论文里，引用可能是编造的——DOI 格式对、但论文不存在"),
    ("张冠李戴", "引用说「A 支持 B」，但原文根本没这么说，甚至说的是反面"),
    ("撤稿论文", "你引的论文可能早就因为数据造假被撤了，但你不知道"),
    ("引用操纵", "有些期刊/作者会互相刷引用，形成「引用圈」"),
    ("数据编造", "统计数据的均值和样本量对不上 (GRIM test)"),
]
for title, desc in problems:
    add_para(tf, "", font_size=6, space_before=8)
    p = add_para(tf, f"  {title}", font_size=16, color=YELLOW, bold=True, space_before=4)
    add_para(tf, f"    {desc}", font_size=13, color=GRAY, space_before=2)

# Right column - scale
add_rounded_rect(slide, 7.0, 1.5, 5.5, 5.2, BG_CARD)
_, tf = add_text_box(slide, 7.2, 1.6, 5.1, 5.0,
                     "这个问题有多严重？", font_size=20, color=ACCENT, bold=True)
stats = [
    "Nature 2023 调查：ChatGPT 生成的引用中\n约 30-70% 是完全虚构的",
    "Retraction Watch 数据库：\n累计 4 万+ 篇论文被撤稿",
    "Cabanac et al. 2021 发现\n数百篇论文含有「扭曲短语」\n(tortured phrases)，疑似论文工厂产物",
    "一篇错误引用可能被后续几十篇论文\n继续传播，形成「错误引用链」",
    "人工逐条检查一篇论文的参考文献\n→ 平均需要 2-3 小时",
]
for s in stats:
    add_para(tf, f"  {s}", font_size=13, color=GRAY, space_before=12)

# ════════════════════════════════════════
# SLIDE 3 — Solution overview (5 layers)
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_DARK)

add_text_box(slide, 0.8, 0.4, 12, 0.8,
             "我们的方案：5 层验证，从简单到深入",
             font_size=36, color=WHITE, bold=True)

layers = [
    ("L0", "这篇论文存在吗？", "查 62 个数据库，看 DOI/标题能不能查到",
     "比如：DOI 格式对但查不到 → 大概率是编的", GREEN),
    ("L1", "引用意图对吗？", "自动判断：这条引用是「支持」「反对」还是「顺便提一嘴」",
     "用 SciBERT 模型分类，F1 = 86%", ACCENT),
    ("L2", "语义对得上吗？", "把你的声明和原文摘要做 NLI 对比",
     "DeBERTa-v3 模型，准确率 93.5%（超过人类一致性 89%）", ACCENT2),
    ("L3", "引用网络正常吗？", "构建引文图谱，检测 7 种异常模式",
     "比如：自引过高、互引圈、时间异常", ORANGE),
    ("L4", "综合风险多高？", "把前面所有信号用贝叶斯公式算出一个风险概率",
     "18 个信号 → 后验概率 → 4 个风险等级", RED),
]

for i, (level, question, method, detail, color) in enumerate(layers):
    y = 1.5 + i * 1.1
    # level badge
    add_rounded_rect(slide, 0.8, y, 0.9, 0.85, color)
    add_text_box(slide, 0.8, y + 0.15, 0.9, 0.5,
                 level, font_size=22, color=WHITE, bold=True,
                 alignment=PP_ALIGN.CENTER)
    # question
    add_text_box(slide, 1.9, y, 3.0, 0.45,
                 question, font_size=18, color=WHITE, bold=True)
    add_text_box(slide, 1.9, y + 0.4, 3.0, 0.45,
                 method, font_size=13, color=GRAY)
    # detail card
    add_rounded_rect(slide, 6.5, y, 6.0, 0.85, BG_CARD)
    add_text_box(slide, 6.7, y + 0.2, 5.6, 0.5,
                 detail, font_size=13, color=GRAY)

# ════════════════════════════════════════
# SLIDE 4 — How L0 works (most intuitive)
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_DARK)

add_text_box(slide, 0.8, 0.4, 12, 0.8,
             "L0 详解：怎么判断一篇引用是不是假的？",
             font_size=36, color=WHITE, bold=True)

_, tf = add_text_box(slide, 0.8, 1.4, 12, 0.6,
                     "思路其实很简单，就像你手动去 Google Scholar 搜一样，只是我们自动化了：",
                     font_size=18, color=GRAY)

steps = [
    ("Step 1: 解析标识符",
     "从参考文献里提取 DOI、arXiv ID、PMID 等——支持 15 种格式自动识别",
     "比如 10.1038/s41586-021-03819-2 → 这是个 Crossref DOI"),
    ("Step 2: 多库查询",
     "拿着 ID 去 62 个数据库同时查：Crossref、OpenAlex、PubMed、DBLP...",
     "查到了 → 真的；查不到 → 可疑；格式对但不存在 → Phantom DOI (很可能是假的)"),
    ("Step 3: 字段比对",
     "就算查到了，也要比对标题、作者、年份是否匹配",
     "用模糊匹配 (RapidFuzz)，容忍大小写、缩写等差异"),
    ("Step 4: 额外检测",
     "查撤稿数据库、扫描扭曲短语 (281 个模式)、GRIM 数值验证",
     "比如 'breast cancer' 被替换成 'bosom disease' → 论文工厂特征"),
]

for i, (title, desc, example) in enumerate(steps):
    y = 2.3 + i * 1.2
    add_rounded_rect(slide, 0.8, y, 11.7, 1.05, BG_CARD)
    add_text_box(slide, 1.0, y + 0.05, 5.0, 0.4,
                 title, font_size=16, color=ACCENT, bold=True)
    add_text_box(slide, 1.0, y + 0.35, 5.5, 0.35,
                 desc, font_size=13, color=WHITE)
    add_text_box(slide, 6.8, y + 0.25, 5.5, 0.55,
                 example, font_size=12, color=YELLOW)

# ════════════════════════════════════════
# SLIDE 5 — L2 Semantic verification
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_DARK)

add_text_box(slide, 0.8, 0.4, 12, 0.8,
             "L2 详解：你的引用真的支持你的说法吗？",
             font_size=36, color=WHITE, bold=True)

_, tf = add_text_box(slide, 0.8, 1.3, 12, 0.6,
                     "这一层做的事情叫 NLI (Natural Language Inference)——判断两句话之间的逻辑关系",
                     font_size=18, color=GRAY)

# Example
add_rounded_rect(slide, 0.8, 2.1, 11.7, 3.8, BG_CARD)
add_text_box(slide, 1.0, 2.2, 11, 0.4,
             "举个例子", font_size=18, color=ACCENT, bold=True)

add_text_box(slide, 1.0, 2.7, 2.5, 0.4,
             "你论文里写的：", font_size=14, color=YELLOW, bold=True)
add_text_box(slide, 3.5, 2.7, 8.5, 0.4,
             "\"Transformer 在翻译任务上超越了 RNN\"  [引用 Vaswani et al. 2017]",
             font_size=14, color=WHITE)

add_text_box(slide, 1.0, 3.3, 2.5, 0.4,
             "原文摘要说的：", font_size=14, color=YELLOW, bold=True)
add_text_box(slide, 3.5, 3.3, 8.5, 0.7,
             "\"Experiments on WMT 2014 show the model achieves 28.4 BLEU, "
             "surpassing the best previously reported RNN-based models.\"",
             font_size=14, color=WHITE)

add_text_box(slide, 1.0, 4.1, 2.5, 0.4,
             "模型判断：", font_size=14, color=YELLOW, bold=True)
add_text_box(slide, 3.5, 4.1, 3.0, 0.4,
             "SUPPORTED (置信度 0.91)", font_size=16, color=GREEN, bold=True)

add_text_box(slide, 1.0, 4.8, 11, 0.8,
             "如果原文说的是反面呢？→ 模型会输出 CONTRADICTED\n"
             "如果原文根本没提这事？→ 模型会输出 NOT_ENOUGH_INFO",
             font_size=14, color=GRAY)

# Tech details
add_rounded_rect(slide, 0.8, 6.1, 5.5, 1.0, BG_CARD)
add_text_box(slide, 1.0, 6.2, 5.1, 0.8,
             "模型：DeBERTa-v3 (微软)\n"
             "在 SciFact 数据集上微调 → 准确率 93.5%\n"
             "人类标注者之间的一致性只有 89.1%",
             font_size=13, color=GRAY)

add_rounded_rect(slide, 6.8, 6.1, 5.7, 1.0, BG_CARD)
add_text_box(slide, 7.0, 6.2, 5.3, 0.8,
             "加速：导出为 ONNX 格式 + INT8 量化\n"
             "CPU 上也能跑，速度提升 2.2 倍\n"
             "不需要 GPU 也能用",
             font_size=13, color=GRAY)

# ════════════════════════════════════════
# SLIDE 6 — Architecture / Tech stack
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_DARK)

add_text_box(slide, 0.8, 0.4, 12, 0.8,
             "技术架构：用了哪些技术？",
             font_size=36, color=WHITE, bold=True)

# Cards grid
cards = [
    ("Python 主体", [
        "核心逻辑全部 Python",
        "aiohttp 异步并发请求",
        "asyncio.gather 并行查询",
        "RapidFuzz 模糊字符串匹配",
    ]),
    ("NLP 模型", [
        "SciBERT — 引文意图分类",
        "DeBERTa-v3 — 语义验证 (NLI)",
        "HuggingFace Transformers",
        "ONNX Runtime 推理加速",
    ]),
    ("数据源 (62 个 API)", [
        "学术: Crossref, OpenAlex, PubMed...",
        "专利: USPTO, EPO, WIPO...",
        "法律: EUR-Lex, CourtListener...",
        "金融/政府/标准 等 6 大领域",
    ]),
    ("工程化", [
        "熔断器: API 挂了自动隔离",
        "双层缓存: LRU + Redis",
        "批处理: 32 并发 Semaphore",
        "Rust 加速: PyO3 绑定 (可选)",
    ]),
    ("评分系统", [
        "贝叶斯后验概率",
        "18 个信号 × 似然比",
        "Kill-shot: 致命信号直接判死刑",
        "4 级风险: LOW → CRITICAL",
    ]),
    ("测试与 Benchmark", [
        "480+ 单元测试",
        "58-case Golden Test Set",
        "与 CheckIfExist 正面对比",
        "消融实验 + 级联早停优化",
    ]),
]

for i, (title, bullets) in enumerate(cards):
    col = i % 3
    row = i // 3
    x = 0.8 + col * 4.1
    y = 1.5 + row * 2.8
    add_card(slide, x, y, 3.8, 2.4, title, bullets, title_color=ACCENT)

# ════════════════════════════════════════
# SLIDE 7 — Results
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_DARK)

add_text_box(slide, 0.8, 0.4, 12, 0.8,
             "实验结果：到底好不好使？", font_size=36, color=WHITE, bold=True)

add_text_box(slide, 0.8, 1.2, 12, 0.5,
             "在 58 个精心设计的测试用例上 (含 14 个假引用、23 篇撤稿、3 个嵌合体、18 篇真论文)：",
             font_size=16, color=GRAY)

# Results table as cards
results = [
    ("假引用检出率", "100%", "14/14 全部抓到", GREEN),
    ("真论文误报率", "0%", "18 篇真论文零误报", GREEN),
    ("撤稿检出率", "73.9%", "17/23 (受限于数据库覆盖)", YELLOW),
    ("嵌合体检测", "100%", "3/3 元数据拼接的假论文", GREEN),
]

for i, (label, value, note, color) in enumerate(results):
    x = 0.8 + i * 3.1
    add_rounded_rect(slide, x, 1.9, 2.8, 1.8, BG_CARD)
    add_text_box(slide, x + 0.1, 1.95, 2.6, 0.4,
                 label, font_size=14, color=GRAY, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + 0.1, 2.4, 2.6, 0.6,
                 value, font_size=36, color=color, bold=True,
                 alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + 0.1, 3.1, 2.6, 0.4,
                 note, font_size=11, color=GRAY, alignment=PP_ALIGN.CENTER)

# vs CheckIfExist
add_text_box(slide, 0.8, 4.0, 12, 0.5,
             "对比现有工具 CheckIfExist (Abbonato et al. 2025)：",
             font_size=18, color=WHITE, bold=True)

add_rounded_rect(slide, 0.8, 4.6, 11.7, 2.5, BG_CARD)
_, tf = add_text_box(slide, 1.0, 4.7, 11.3, 2.3, "", font_size=14)
comparisons = [
    ("假引用检出 (高置信)",   "IntegriRef 100%  vs  CheckIfExist 0%",
     "— CIE 无法给出高置信判断"),
    ("撤稿检出",             "IntegriRef 73.9%  vs  CheckIfExist 26.1%",
     "— CIE 只查存在性，不查撤稿状态"),
    ("嵌合体检测",           "IntegriRef 100%  vs  CheckIfExist 33.3%",
     "— CIE 对标题正确的嵌合体完全盲目"),
    ("误报率",              "IntegriRef 0%  vs  CheckIfExist 11.1%",
     "— CIE 因标题格式差异误判"),
    ("速度",                "IntegriRef 38.5s  vs  CheckIfExist 2.3s",
     "— 我们慢但做了 5 层验证，CIE 只查存在性"),
]
for metric, vals, note in comparisons:
    add_para(tf, f"  {metric}:  {vals} {note}",
             font_size=13, color=GRAY, space_before=6)

# ════════════════════════════════════════
# SLIDE 8 — Demo / Code example
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_DARK)

add_text_box(slide, 0.8, 0.4, 12, 0.8,
             "怎么用？三行代码就能跑",
             font_size=36, color=WHITE, bold=True)

code1 = '''from core.discovery import RegistryDiscovery

discovery = RegistryDiscovery()
discovery.auto_register()   # 自动注册 62 个数据库适配器

result = discovery.verify_reference(
    title="Attention Is All You Need",
    identifier="10.48550/arXiv.1706.03762",
)
print(result["found"])            # True
print(result["registries_hit"])   # ["crossref", "openalex", ...]'''

add_rounded_rect(slide, 0.8, 1.4, 5.5, 3.6, BG_CARD)
add_text_box(slide, 1.0, 1.5, 5.1, 0.3,
             "基本用法：查一条引用是否存在", font_size=14, color=ACCENT, bold=True)
add_text_box(slide, 1.0, 1.9, 5.1, 3.0,
             code1, font_size=11, color=GREEN, font_name="Consolas")

code2 = '''from semantic.nli_verifier import NLIVerifier

verifier = NLIVerifier()  # 加载 DeBERTa-v3
result = verifier.verify(
    claim="Transformer 超越了 RNN",
    abstract="...achieves 28.4 BLEU, "
             "surpassing RNN-based models."
)
# result.label → SUPPORTED
# result.confidence → 0.91'''

add_rounded_rect(slide, 7.0, 1.4, 5.5, 3.6, BG_CARD)
add_text_box(slide, 7.2, 1.5, 5.1, 0.3,
             "语义验证：声明和原文对得上吗？", font_size=14, color=ACCENT, bold=True)
add_text_box(slide, 7.2, 1.9, 5.1, 3.0,
             code2, font_size=11, color=GREEN, font_name="Consolas")

# Bottom - pipeline
add_rounded_rect(slide, 0.8, 5.3, 11.7, 1.8, BG_CARD)
add_text_box(slide, 1.0, 5.4, 11, 0.3,
             "完整管线：一键跑 L0 → L4", font_size=14, color=ACCENT, bold=True)

code3 = '''from core.pipeline import IntegriRefPipeline

pipeline = IntegriRefPipeline()
report = pipeline.verify(references)      # 输入一组参考文献
# report 包含：每条引用的风险等级 (LOW/ELEVATED/HIGH/CRITICAL)、
#              具体触发了哪些信号、以及贝叶斯后验概率'''

add_text_box(slide, 1.0, 5.8, 11.3, 1.2,
             code3, font_size=11, color=GREEN, font_name="Consolas")

# ════════════════════════════════════════
# SLIDE 9 — What's next
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_DARK)

add_text_box(slide, 0.8, 0.4, 12, 0.8,
             "接下来要做什么？", font_size=36, color=WHITE, bold=True)

# TODO items
todos = [
    ("正在做", ACCENT, [
        ("系统论文撰写", "投 JASIST/Scientometrics，L0-L4 全栈 + 消融实验"),
        ("数据集扩展", "用爬虫收集更多撤稿论文、统计错误、幻觉引用样本"),
        ("接入剩余检测器", "4 个检测器 (扭曲短语/邮箱/GRIM/Sneaked) 还没接进管线"),
    ]),
    ("计划中", YELLOW, [
        ("LLM 幻觉引文检测", "专门检测 GPT-4/Claude/Gemini 生成的假引用 (投 ACL)"),
        ("REST API", "FastAPI 封装，让其他工具和编辑器可以直接调用"),
        ("跨数据集泛化", "在 HealthVer、FEVER、ClimateVER 上测试 L2 的泛化能力"),
    ]),
    ("远期", GRAY, [
        ("浏览器插件", "在 Google Scholar / arXiv 页面直接显示风险评分"),
        ("多语言优化", "针对中文、日文等非英语论文的验证精度提升"),
        ("实时监控", "监控新发表论文的引用质量，自动预警"),
    ]),
]

y = 1.4
for section, color, items in todos:
    add_text_box(slide, 0.8, y, 2.0, 0.4,
                 section, font_size=20, color=color, bold=True)
    y += 0.45
    for title, desc in items:
        add_rounded_rect(slide, 0.8, y, 11.7, 0.55, BG_CARD)
        add_text_box(slide, 1.0, y + 0.05, 3.5, 0.4,
                     title, font_size=14, color=WHITE, bold=True)
        add_text_box(slide, 4.5, y + 0.05, 7.8, 0.4,
                     desc, font_size=13, color=GRAY)
        y += 0.62
    y += 0.15

# ════════════════════════════════════════
# SLIDE 10 — Thank you
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_DARK)

add_text_box(slide, 1.5, 2.0, 10, 1.0,
             "Thanks!", font_size=54, color=ACCENT, bold=True,
             alignment=PP_ALIGN.CENTER)

add_text_box(slide, 1.5, 3.2, 10, 0.6,
             "一句话总结：帮你自动检查参考文献靠不靠谱的开源工具",
             font_size=20, color=WHITE, alignment=PP_ALIGN.CENTER)

_, tf = add_text_box(slide, 1.5, 4.2, 10, 2.5, "", font_size=16,
                     alignment=PP_ALIGN.CENTER)
add_para(tf, "GitHub: GeoffreyWang1117/IntegriRef",
         font_size=16, color=GRAY, space_before=12, alignment=PP_ALIGN.CENTER)
add_para(tf, "开源协议: MIT License",
         font_size=16, color=GRAY, space_before=8, alignment=PP_ALIGN.CENTER)
add_para(tf, "", font_size=10, space_before=20)
add_para(tf, "欢迎 Star / Fork / Issue / PR",
         font_size=18, color=ACCENT, space_before=8, alignment=PP_ALIGN.CENTER)

# ── Save ──
output_path = "/home/coder-gw/Projects/IntegriRef/IntegriRef_Intro.pptx"
prs.save(output_path)
print(f"Saved to {output_path}")
